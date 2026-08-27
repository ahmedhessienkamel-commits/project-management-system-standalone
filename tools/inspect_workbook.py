from pathlib import Path
import json
import openpyxl

path = Path('/home/ubuntu/upload/ممم.xlsx')
workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
result = []
for sheet in workbook.worksheets:
    rows = sheet.iter_rows(values_only=True)
    sample = []
    for index, row in enumerate(rows):
        values = list(row[:12])
        if any(value not in (None, '') for value in values):
            sample.append(values)
        if len(sample) >= 5:
            break
    result.append({'title': sheet.title, 'max_row': sheet.max_row, 'max_column': sheet.max_column, 'sample': sample})
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
